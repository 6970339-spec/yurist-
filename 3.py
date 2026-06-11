import os
import csv
import json
import sqlite3
import platform
import subprocess
import importlib
import re
from copy import deepcopy
import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry
from docxtpl import DocxTemplate
import pymorphy3

from banks import (
    BanksDatabaseError,
    add_bank,
    delete_bank,
    get_bank_details,
    search_banks,
    update_bank,
)

FONT = ("Times New Roman", 12)
BG_OK = "white"
BG_ERROR = "#ffcccc"

TEMPLATES = {
    "zayavlenie": "templates/zayavlenie.docx",
    "creditors": "templates/creditors.docx",
    "property": "templates/property.docx",
}

OUTPUT_DIR = "output"
DATA_DIR = "data"
CLIENTS_FILE = os.path.join(DATA_DIR, "clients.json")

CODES_TXT_FILE = "codes_pass.txt"
CODES_DB_FILE = os.path.join(DATA_DIR, "passport_codes.db")
DADATA_TOKEN = "cae6427b82a9aa8e2289ffc19e06332c9c6d08ee"

morph = pymorphy3.MorphAnalyzer()

fields = [
    ("Фамилия", "surname"),
    ("Имя", "name"),
    ("Отчество", "patronymic"),
    ("ФИО в родительном падеже", "fio_genitive"),
    ("Дата рождения", "birth_date"),
    ("Место рождения", "birth_place"),
    ("СНИЛС", "snils"),
    ("ИНН", "inn"),
    ("Серия паспорта", "passport_series"),
    ("Номер паспорта", "passport_number"),
    ("Дата выдачи паспорта", "passport_issue_date"),
    ("Код подразделения", "passport_code"),
    ("Кем выдан паспорт", "passport_issued_by"),
    ("Индекс", "postal_code"),
    ("Регион", "region"),
    ("Район", "district"),
    ("Город", "city"),
    ("Населённый пункт", "locality"),
    ("Улица", "street"),
    ("Дом", "house"),
    ("Корпус", "building"),
    ("Квартира", "apartment"),
    ("Доход", "income"),
]

manual_date_keys = {"birth_date", "passport_issue_date"}

digit_limits = {
    "snils": 11,
    "inn": 12,
    "passport_series": 4,
    "passport_number": 6,
    "postal_code": 6,
}

month_names = {
    "01": "января", "02": "февраля", "03": "марта",
    "04": "апреля", "05": "мая", "06": "июня",
    "07": "июля", "08": "августа", "09": "сентября",
    "10": "октября", "11": "ноября", "12": "декабря",
}


def ensure_files():
    os.makedirs(DATA_DIR, exist_ok=True)
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    if not os.path.exists(CLIENTS_FILE):
        with open(CLIENTS_FILE, "w", encoding="utf-8") as file:
            json.dump([], file, ensure_ascii=False, indent=4)


def only_digits(value, max_len=None):
    digits = "".join(ch for ch in value if ch.isdigit())
    return digits[:max_len] if max_len else digits


def format_passport_code(value):
    digits = only_digits(value, 6)
    if len(digits) > 3:
        return f"{digits[:3]}-{digits[3:]}"
    return digits


def normalize_passport_code(value):
    digits = only_digits(value, 6)
    if len(digits) == 6:
        return f"{digits[:3]}-{digits[3:]}"
    return value.strip()


def normalize_gosuslugi_label(value):
    return re.sub(r"[^а-яёa-z0-9]+", "", value.lower())


def find_next_gosuslugi_value(lines, labels):
    normalized_labels = {normalize_gosuslugi_label(label) for label in labels}

    for index, line in enumerate(lines):
        if normalize_gosuslugi_label(line) in normalized_labels:
            for next_line in lines[index + 1:]:
                if next_line.strip():
                    return next_line.strip()

    return ""


def find_value_after_label_in_text(text, label, pattern):
    match = re.search(
        rf"{label}\s*[:\n\r]+(?P<value>[\s\S]{{0,120}}?)({pattern})",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(2) if match else ""


def parse_gosuslugi_address(address_text):
    result = {}
    text = " ".join(address_text.replace("\n", ", ").split())

    postal_match = re.search(r"\b(\d{6})\b", text)
    if postal_match:
        result["postal_code"] = postal_match.group(1)

    parts = [part.strip() for part in re.split(r"[,;]", text) if part.strip()]

    for part in parts:
        lower_part = part.lower()

        if "район" in lower_part or " р-н" in lower_part:
            result.setdefault("district", part)
        elif any(word in lower_part for word in ("республика", "область", "край", "автоном")):
            result.setdefault("region", part)
        elif re.search(r"\bг\.?\s+", lower_part) or "город" in lower_part:
            result.setdefault("city", part)
        elif any(word in lower_part for word in ("село", "поселок", "посёлок", "деревня", "пгт")) or re.search(r"\b[сдп]\.?\s+", lower_part):
            result.setdefault("locality", part)
        elif re.search(r"\bул\.?\s+", lower_part) or "улица" in lower_part:
            result.setdefault("street", part)

    house_match = re.search(r"(?:\bд\.?|\bдом)\s*([\wА-Яа-яёЁ/-]+)", text, flags=re.IGNORECASE)
    if house_match:
        result["house"] = house_match.group(1)

    building_match = re.search(r"(?:корп\.?|корпус)\s*([\wА-Яа-яёЁ/-]+)", text, flags=re.IGNORECASE)
    if building_match:
        result["building"] = building_match.group(1)

    apartment_match = re.search(r"(?:кв\.?|квартира)\s*([\wА-Яа-яёЁ/-]+)", text, flags=re.IGNORECASE)
    if apartment_match:
        result["apartment"] = apartment_match.group(1)

    return result


def collect_gosuslugi_address_text(lines):
    address_labels = {"адресрегистрации", "адрес", "регистрация"}
    stop_labels = {
        "паспорт", "выдан", "кодподразделения", "датавыдачи", "снилс", "инн",
        "фамилия", "имя", "отчество", "датарождения", "месторождения",
    }

    for index, line in enumerate(lines):
        normalized = normalize_gosuslugi_label(line)
        if normalized in address_labels:
            address_lines = []
            for next_line in lines[index + 1:]:
                normalized_next = normalize_gosuslugi_label(next_line)
                if normalized_next in stop_labels:
                    break
                if next_line.strip():
                    address_lines.append(next_line.strip())
            return " ".join(address_lines)

    return ""


def parse_gosuslugi_text(text):
    result = {}
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    passport_value = find_next_gosuslugi_value(lines, ["Паспорт"])
    passport_digits = only_digits(passport_value)
    if len(passport_digits) >= 10:
        result["passport_series"] = passport_digits[:4]
        result["passport_number"] = passport_digits[4:10]
    else:
        passport_match = re.search(r"\b(\d{4})\s+(\d{6})\b", text)
        if passport_match:
            result["passport_series"] = passport_match.group(1)
            result["passport_number"] = passport_match.group(2)

    passport_issued_by = find_next_gosuslugi_value(lines, ["Выдан", "Кем выдан"])
    if passport_issued_by:
        result["passport_issued_by"] = passport_issued_by

    passport_code = find_next_gosuslugi_value(lines, ["Код подразделения"])
    if passport_code:
        result["passport_code"] = normalize_passport_code(passport_code)

    passport_issue_date = find_next_gosuslugi_value(lines, ["Дата выдачи"])
    if passport_issue_date:
        result["passport_issue_date"] = passport_issue_date

    snils_text = find_next_gosuslugi_value(lines, ["СНИЛС"])
    snils_digits = only_digits(snils_text) or only_digits(find_value_after_label_in_text(text, "СНИЛС", r"\d{3}[-\s]?\d{3}[-\s]?\d{3}\s?\d{2}"))
    if len(snils_digits) >= 11:
        result["snils"] = snils_digits[:11]

    inn_text = find_next_gosuslugi_value(lines, ["ИНН"])
    inn_digits = only_digits(inn_text) or only_digits(find_value_after_label_in_text(text, "ИНН", r"\d{12}"))
    if len(inn_digits) >= 12:
        result["inn"] = inn_digits[:12]

    for labels, key in ((["Фамилия"], "surname"), (["Имя"], "name"), (["Отчество"], "patronymic")):
        value = find_next_gosuslugi_value(lines, labels)
        if value:
            result[key] = value

    birth_date = find_next_gosuslugi_value(lines, ["Дата рождения"])
    if birth_date:
        result["birth_date"] = birth_date

    birth_place = find_next_gosuslugi_value(lines, ["Место рождения"])
    if birth_place:
        result["birth_place"] = birth_place

    address_text = collect_gosuslugi_address_text(lines)
    if address_text:
        result.update(parse_gosuslugi_address(address_text))

    return result


def load_clients():
    ensure_files()
    try:
        with open(CLIENTS_FILE, "r", encoding="utf-8") as file:
            return json.load(file)
    except Exception:
        return []


def save_clients(clients):
    ensure_files()
    with open(CLIENTS_FILE, "w", encoding="utf-8") as file:
        json.dump(clients, file, ensure_ascii=False, indent=4)


def recreate_passport_codes_db():
    ensure_files()

    if not os.path.exists(CODES_TXT_FILE):
        return

    if os.path.exists(CODES_DB_FILE):
        os.remove(CODES_DB_FILE)

    conn = sqlite3.connect(CODES_DB_FILE)
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE passport_codes (
            id TEXT,
            name TEXT NOT NULL,
            code TEXT NOT NULL UNIQUE,
            end_date TEXT
        )
    """)

    cursor.execute("CREATE INDEX idx_passport_codes_code ON passport_codes(code)")
    cursor.execute("CREATE INDEX idx_passport_codes_name ON passport_codes(name)")

    encodings = ["utf-8-sig", "utf-8", "cp1251"]
    opened_file = None

    for encoding in encodings:
        try:
            opened_file = open(CODES_TXT_FILE, "r", encoding=encoding, newline="")
            opened_file.read(100)
            opened_file.seek(0)
            break
        except UnicodeDecodeError:
            if opened_file:
                opened_file.close()
            opened_file = None

    if not opened_file:
        conn.close()
        return

    rows = []

    with opened_file as file:
        reader = csv.DictReader(file, delimiter=";")
        for row in reader:
            code = normalize_passport_code(row.get("CODE", ""))
            name = row.get("NAME", "").strip()
            record_id = row.get("ID", "").strip()
            end_date = row.get("END_DATE", "").strip()

            if code and name and len(only_digits(code)) == 6:
                rows.append((record_id, name, code, end_date))

    cursor.executemany("""
        INSERT OR REPLACE INTO passport_codes (id, name, code, end_date)
        VALUES (?, ?, ?, ?)
    """, rows)

    conn.commit()
    conn.close()


def create_passport_codes_db_if_needed():
    if not os.path.exists(CODES_DB_FILE):
        recreate_passport_codes_db()


def search_address_by_postal_code(postal_code):
    try:
        dadata_module = importlib.import_module("dadata")
    except ImportError as exc:
        raise RuntimeError("Установите библиотеку: pip install dadata") from exc

    try:
        dadata = dadata_module.Dadata(DADATA_TOKEN)
        result = dadata.suggest(
            "address",
            "",
            count=20,
            locations=[{"postal_code": postal_code}],
        )
        print("DADATA RESULT:", result)
        return result
    except Exception as exc:
        print("DADATA ERROR:", repr(exc))
        raise RuntimeError(f"Не удалось получить адрес по индексу: {exc}") from exc
    finally:
        close = locals().get("dadata") and getattr(dadata, "close", None)
        if close:
            close()


def search_passport_codes(query, limit=10):
    create_passport_codes_db_if_needed()

    if not os.path.exists(CODES_DB_FILE):
        return []

    query = query.strip()
    digits = only_digits(query)

    conn = sqlite3.connect(CODES_DB_FILE)
    cursor = conn.cursor()

    if digits:
        like_value = f"{digits}%"
        cursor.execute("""
            SELECT code, name
            FROM passport_codes
            WHERE REPLACE(code, '-', '') LIKE ?
            ORDER BY code
            LIMIT ?
        """, (like_value, limit))
    else:
        like_value = f"%{query.upper()}%"
        cursor.execute("""
            SELECT code, name
            FROM passport_codes
            WHERE UPPER(name) LIKE ?
            ORDER BY code
            LIMIT ?
        """, (like_value, limit))

    rows = cursor.fetchall()
    conn.close()
    return rows


def count_codes_in_db():
    create_passport_codes_db_if_needed()

    if not os.path.exists(CODES_DB_FILE):
        return 0

    conn = sqlite3.connect(CODES_DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM passport_codes")
    count = cursor.fetchone()[0]
    conn.close()
    return count


def auto_date_format(event):
    entry = event.widget
    digits = only_digits(entry.get(), 8)

    result = ""
    for i, ch in enumerate(digits):
        if i == 2 or i == 4:
            result += "."
        result += ch

    entry.delete(0, tk.END)
    entry.insert(0, result)
    entry.configure(bg=BG_OK if len(result) == 10 else BG_ERROR)


def auto_digits_format(event, max_len):
    entry = event.widget
    digits = only_digits(entry.get(), max_len)
    entry.delete(0, tk.END)
    entry.insert(0, digits)
    entry.configure(bg=BG_OK if len(digits) == max_len else BG_ERROR)


def to_genitive_word(word):
    if not word:
        return ""

    parsed = morph.parse(word)[0]
    genitive = parsed.inflect({"gent"})

    return genitive.word.capitalize() if genitive else word


def fio_to_genitive(surname, name, patronymic):
    return " ".join(
        to_genitive_word(part)
        for part in [surname, name, patronymic]
        if part
    )


def split_date(date_text, prefix):
    result = {
        f"{prefix}_day": "",
        f"{prefix}_month_number": "",
        f"{prefix}_month": "",
        f"{prefix}_year": "",
    }

    parts = date_text.split(".")
    if len(parts) == 3:
        day, month, year = parts
        result[f"{prefix}_day"] = day
        result[f"{prefix}_month_number"] = month
        result[f"{prefix}_month"] = month_names.get(month, "")
        result[f"{prefix}_year"] = year

    return result


def safe_folder_name(name):
    bad_chars = '<>:"/\\|?*'
    for char in bad_chars:
        name = name.replace(char, "_")
    return name.strip()


def open_folder(path):
    system_name = platform.system()

    if system_name == "Windows":
        os.startfile(path)
    elif system_name == "Darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


def parse_money(value):
    text = str(value or "").replace(" ", "").replace(" ", "").replace(",", ".")
    allowed = "".join(ch for ch in text if ch.isdigit() or ch == ".")

    if not allowed:
        return 0.0

    try:
        return float(allowed)
    except ValueError:
        return 0.0


def format_money(value):
    amount = parse_money(value)

    if amount.is_integer():
        return f"{int(amount):,}".replace(",", " ")

    return f"{amount:,.2f}".replace(",", " ").replace(".", ",")


def prepare_creditors_for_documents(creditors):
    prepared_creditors = []
    total_debt = 0.0

    for index, creditor in enumerate(creditors, start=1):
        contract_date = creditor.get("contract_date", "")
        prepared_creditor = {
            "row_number": f"1.{index}",
            "obligation_type": "Кредит",
            "name": creditor.get("name", ""),
            "inn": creditor.get("inn", ""),
            "ogrn": creditor.get("ogrn", ""),
            "address": creditor.get("address", ""),
            "contract_basis": f"Кредитный договор от {contract_date}".strip(),
            "debt_sum": creditor.get("debt_sum", ""),
            "contract_date": contract_date,
            "contract_number": creditor.get("contract_number", ""),
            "penalty": "0",
        }
        debt_value = parse_money(prepared_creditor["debt_sum"])
        total_debt += debt_value
        prepared_creditor["debt_sum_formatted"] = format_money(debt_value)
        prepared_creditors.append(prepared_creditor)

    return prepared_creditors, total_debt


def get_creditors_template_render_data(data):
    render_data = data.copy()
    render_data["creditors"] = [{
        "row_number": "__CREDITOR_TEMPLATE_ROW__",
        "obligation_type": "__CREDITOR_TEMPLATE_ROW__",
        "name": "__CREDITOR_TEMPLATE_ROW__",
        "address": "__CREDITOR_TEMPLATE_ROW__",
        "inn": "__CREDITOR_TEMPLATE_ROW__",
        "ogrn": "__CREDITOR_TEMPLATE_ROW__",
        "contract_basis": "__CREDITOR_TEMPLATE_ROW__",
        "debt_sum": "__CREDITOR_TEMPLATE_ROW__",
        "debt_sum_formatted": "__CREDITOR_TEMPLATE_ROW__",
        "contract_date": "__CREDITOR_TEMPLATE_ROW__",
        "contract_number": "__CREDITOR_TEMPLATE_ROW__",
        "penalty": "__CREDITOR_TEMPLATE_ROW__",
    }]
    render_data["creditor"] = render_data["creditors"][0]
    return render_data


def row_text(row):
    return "\n".join(cell.text for cell in row.cells)


def find_creditors_table(document):
    for table in document.tables:
        table_text = "\n".join(row_text(row) for row in table.rows)
        if "Сведения о кредиторах гражданина" in table_text and "Кредитор" in table_text:
            return table

    for table in document.tables:
        table_text = "\n".join(row_text(row) for row in table.rows)
        if "__CREDITOR_TEMPLATE_ROW__" in table_text or "creditor." in table_text or "{%tr" in table_text:
            return table

    return None


def find_creditor_template_row(table):
    for row in table.rows:
        text = row_text(row)
        if "__CREDITOR_TEMPLATE_ROW__" in text:
            return row
        if "creditor." in text or "{%tr" in text or "loop.index" in text:
            return row

    for row in table.rows:
        cells = row.cells
        if len(cells) >= 8 and cells[0].text.strip().startswith("1."):
            return row

    return None


def clear_cell(cell):
    for paragraph in cell.paragraphs:
        paragraph.text = ""


def set_cell_text(cell, text):
    clear_cell(cell)

    if cell.paragraphs:
        cell.paragraphs[0].text = str(text or "")
    else:
        cell.text = str(text or "")


def fill_creditor_row(row, creditor):
    values = (
        creditor.get("row_number", ""),
        creditor.get("obligation_type", "Кредит"),
        creditor.get("name", ""),
        creditor.get("address", ""),
        creditor.get("contract_basis", ""),
        f"{creditor.get('debt_sum_formatted', '')} руб.",
        f"{creditor.get('debt_sum_formatted', '')} руб.",
        creditor.get("penalty", "0"),
    )

    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def render_creditors_table_with_python_docx(output_path, creditors):
    from docx import Document
    from docx.table import _Row

    document = Document(output_path)
    table = find_creditors_table(document)

    if table is None:
        raise ValueError("Не найдена таблица 'Сведения о кредиторах гражданина'.")

    template_row = find_creditor_template_row(table)
    if template_row is None:
        raise ValueError("Не найдена строка-шаблон кредитора в creditors.docx.")

    template_tr = template_row._tr
    insert_after = template_tr

    for creditor in creditors:
        new_tr = deepcopy(template_tr)
        insert_after.addnext(new_tr)
        insert_after = new_tr
        fill_creditor_row(_Row(new_tr, table), creditor)

    template_tr.getparent().remove(template_tr)
    document.save(output_path)


class YuristApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Yurist+")
        self.root.geometry("980x860")

        self.entries = {}
        self.clients = load_clients()
        self.codes_count = count_codes_in_db()
        self.current_code_results = []
        self.bank_window = None
        self.creditors = []
        self.last_postal_code_lookup = ""
        self.postal_code_after_id = None

        self.create_menu()
        self.create_ui()
        self.refresh_clients_list()

    def create_menu(self):
        menu_bar = tk.Menu(self.root)

        file_menu = tk.Menu(menu_bar, tearoff=0)
        file_menu.add_command(label="Выход", command=self.root.quit)
        menu_bar.add_cascade(label="Файл", menu=file_menu)

        menu_bar.add_command(label="Банки", command=self.open_banks_window)

        self.root.config(menu=menu_bar)

    def open_banks_window(self):
        if self.bank_window and self.bank_window.window.winfo_exists():
            self.bank_window.window.lift()
            self.bank_window.window.focus_force()
            return

        self.bank_window = BanksWindow(self.root)

    def create_ui(self):
        canvas = tk.Canvas(self.root)
        scrollbar = ttk.Scrollbar(self.root, orient="vertical", command=canvas.yview)
        frame = ttk.Frame(canvas)

        frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))

        canvas.create_window((0, 0), window=frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        tk.Label(
            frame,
            text="Yurist+ — карточка клиента",
            font=("Times New Roman", 16, "bold")
        ).grid(row=0, column=0, columnspan=3, pady=15)

        row = 1

        tk.Label(frame, text="Сохранённые клиенты", font=FONT).grid(
            row=row, column=0, sticky="w", padx=10, pady=4
        )

        self.client_combo = ttk.Combobox(frame, width=52, font=FONT, state="readonly")
        self.client_combo.grid(row=row, column=1, padx=10, pady=4)

        tk.Button(
            frame,
            text="Загрузить",
            font=FONT,
            command=self.load_selected_client
        ).grid(row=row, column=2, padx=10, pady=4)

        row += 1

        tk.Label(
            frame,
            text=f"База кодов DB: {self.codes_count} записей",
            font=FONT
        ).grid(row=row, column=0, columnspan=3, sticky="w", padx=10, pady=4)

        row += 1

        tk.Button(
            frame,
            text="Вставить данные с Госуслуг",
            font=FONT,
            command=self.open_gosuslugi_import_dialog,
            width=30,
        ).grid(row=row, column=0, columnspan=3, pady=8)

        row += 1

        for label_text, key in fields:
            tk.Label(frame, text=label_text, font=FONT).grid(
                row=row, column=0, sticky="w", padx=10, pady=4
            )

            entry = tk.Entry(frame, width=55, font=FONT)
            entry.grid(row=row, column=1, padx=10, pady=4)

            if key in manual_date_keys:
                entry.bind("<KeyRelease>", auto_date_format)

            elif key == "passport_code":
                entry.bind("<KeyRelease>", self.on_passport_code_change)

            elif key == "postal_code":
                entry.bind("<KeyRelease>", self.on_postal_code_change)
                entry.bind("<FocusOut>", self.on_postal_code_focus_out)

            elif key in digit_limits:
                max_len = digit_limits[key]
                entry.bind(
                    "<KeyRelease>",
                    lambda event, limit=max_len: auto_digits_format(event, limit)
                )

            self.entries[key] = entry

            row += 1

            if key == "passport_code":
                self.code_results_box = tk.Listbox(frame, width=90, height=5, font=FONT)
                self.code_results_box.grid(row=row, column=0, columnspan=3, padx=10, pady=4)
                self.code_results_box.bind("<<ListboxSelect>>", self.select_passport_code)
                row += 1

        tk.Button(
            frame,
            text="Авто ФИО в родительном падеже",
            font=FONT,
            command=self.auto_fill_genitive
        ).grid(row=6, column=2, padx=10, pady=4)

        tk.Label(frame, text="Дата заполнения", font=FONT).grid(
            row=row, column=0, sticky="w", padx=10, pady=4
        )

        self.fill_date_entry = DateEntry(
            frame,
            width=20,
            font=FONT,
            date_pattern="dd.mm.yyyy",
            locale="ru_RU"
        )
        self.fill_date_entry.grid(row=row, column=1, sticky="w", padx=10, pady=4)

        row += 1

        row = self.create_creditors_section(frame, row)

        tk.Button(
            frame,
            text="Сохранить клиента в базу",
            font=("Times New Roman", 12, "bold"),
            command=self.save_client_to_base,
            width=30
        ).grid(row=row, column=0, columnspan=3, pady=10)

        row += 1

        tk.Button(
            frame,
            text="Сформировать 3 документа",
            font=("Times New Roman", 12, "bold"),
            command=self.generate_documents,
            width=30
        ).grid(row=row, column=0, columnspan=3, pady=15)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

    def open_gosuslugi_import_dialog(self):
        GosuslugiImportDialog(self.root, self.import_gosuslugi_data)

    def import_gosuslugi_data(self, text):
        parsed_data = parse_gosuslugi_text(text)
        found_fields = []

        for key, value in parsed_data.items():
            if not value or key not in self.entries:
                continue

            self.entries[key].delete(0, tk.END)
            self.entries[key].insert(0, value)
            self.entries[key].configure(bg=BG_OK)
            found_fields.append(key)

        if all(self.entries[key].get().strip() for key in ("surname", "name", "patronymic")):
            self.auto_fill_genitive()

        return found_fields

    def create_creditors_section(self, frame, row):
        tk.Label(
            frame,
            text="Кредиторы / банки",
            font=("Times New Roman", 14, "bold"),
        ).grid(row=row, column=0, columnspan=3, sticky="w", padx=10, pady=(16, 6))
        row += 1

        buttons_frame = ttk.Frame(frame)
        buttons_frame.grid(row=row, column=0, columnspan=3, sticky="w", padx=10, pady=4)

        tk.Button(
            buttons_frame,
            text="Добавить кредитора",
            font=FONT,
            command=self.add_creditor,
            width=20,
        ).pack(side="left", padx=(0, 6))

        tk.Button(
            buttons_frame,
            text="Изменить",
            font=FONT,
            command=self.edit_creditor,
            width=14,
        ).pack(side="left", padx=6)

        tk.Button(
            buttons_frame,
            text="Удалить",
            font=FONT,
            command=self.delete_creditor,
            width=14,
        ).pack(side="left", padx=6)

        row += 1

        columns = (
            "number",
            "name",
            "inn",
            "ogrn",
            "address",
            "debt_sum",
            "contract_date",
            "contract_number",
        )
        self.creditors_tree = ttk.Treeview(
            frame,
            columns=columns,
            show="headings",
            height=8,
            selectmode="browse",
        )

        headings = {
            "number": "№",
            "name": "Банк / кредитор",
            "inn": "ИНН",
            "ogrn": "ОГРН",
            "address": "Адрес",
            "debt_sum": "Сумма долга",
            "contract_date": "Дата договора",
            "contract_number": "Номер договора / идентификатор договора",
        }
        widths = {
            "number": 45,
            "name": 180,
            "inn": 110,
            "ogrn": 130,
            "address": 230,
            "debt_sum": 110,
            "contract_date": 120,
            "contract_number": 220,
        }

        for column in columns:
            self.creditors_tree.heading(column, text=headings[column])
            self.creditors_tree.column(column, width=widths[column], minwidth=45, stretch=True)

        self.creditors_tree.grid(row=row, column=0, columnspan=3, sticky="nsew", padx=10, pady=4)
        self.creditors_tree.bind("<Double-Button-1>", lambda event: self.edit_creditor())

        row += 1

        self.creditors_total_label = tk.Label(
            frame,
            text="Общая сумма задолженности: 0 руб.",
            font=("Times New Roman", 12, "bold"),
        )
        self.creditors_total_label.grid(
            row=row,
            column=0,
            columnspan=3,
            sticky="w",
            padx=10,
            pady=(2, 8),
        )

        row += 1
        return row

    def refresh_creditors_table(self):
        for item_id in self.creditors_tree.get_children():
            self.creditors_tree.delete(item_id)

        for index, creditor in enumerate(self.creditors, start=1):
            self.creditors_tree.insert(
                "",
                tk.END,
                iid=str(index - 1),
                values=(
                    index,
                    creditor.get("name", ""),
                    creditor.get("inn", ""),
                    creditor.get("ogrn", ""),
                    creditor.get("address", ""),
                    creditor.get("debt_sum", ""),
                    creditor.get("contract_date", ""),
                    creditor.get("contract_number", ""),
                ),
            )

        self.update_creditors_total_label()

    def update_creditors_total_label(self):
        if not hasattr(self, "creditors_total_label"):
            return

        _, total_debt = prepare_creditors_for_documents(self.creditors)
        self.creditors_total_label.configure(
            text=f"Общая сумма задолженности: {format_money(total_debt)} руб."
        )

    def get_selected_creditor_index(self):
        selection = self.creditors_tree.selection()
        if not selection:
            return None

        return int(selection[0])

    def add_creditor(self):
        CreditorDialog(
            self.root,
            title="Добавить кредитора",
            on_save=self.save_new_creditor,
        )

    def edit_creditor(self):
        index = self.get_selected_creditor_index()
        if index is None:
            messagebox.showerror("Кредиторы", "Выберите кредитора для изменения.")
            return

        CreditorDialog(
            self.root,
            title="Изменить кредитора",
            creditor=self.creditors[index],
            on_save=lambda data: self.save_existing_creditor(index, data),
        )

    def delete_creditor(self):
        index = self.get_selected_creditor_index()
        if index is None:
            messagebox.showerror("Кредиторы", "Выберите кредитора для удаления.")
            return

        creditor_name = self.creditors[index].get("name", "")
        if not messagebox.askyesno(
            "Удаление кредитора",
            f"Удалить кредитора «{creditor_name}» из карточки клиента?",
        ):
            return

        self.creditors.pop(index)
        self.refresh_creditors_table()

    def save_new_creditor(self, creditor):
        self.creditors.append(creditor)
        self.refresh_creditors_table()
        return True

    def save_existing_creditor(self, index, creditor):
        self.creditors[index] = creditor
        self.refresh_creditors_table()
        return True

    def on_postal_code_change(self, event):
        entry = event.widget
        digits = only_digits(entry.get(), digit_limits["postal_code"])
        entry.delete(0, tk.END)
        entry.insert(0, digits)
        entry.configure(bg=BG_OK if len(digits) == digit_limits["postal_code"] else BG_ERROR)

        if len(digits) == digit_limits["postal_code"]:
            self.schedule_postal_code_lookup(digits)

    def on_postal_code_focus_out(self, event):
        digits = only_digits(event.widget.get(), digit_limits["postal_code"])
        if len(digits) == digit_limits["postal_code"]:
            self.schedule_postal_code_lookup(digits, delay_ms=0)

    def schedule_postal_code_lookup(self, postal_code, delay_ms=600):
        if postal_code == self.last_postal_code_lookup:
            return

        if self.postal_code_after_id:
            self.root.after_cancel(self.postal_code_after_id)

        self.postal_code_after_id = self.root.after(
            delay_ms,
            lambda code=postal_code: self.lookup_postal_code(code),
        )

    def lookup_postal_code(self, postal_code):
        self.postal_code_after_id = None

        if postal_code == self.last_postal_code_lookup:
            return

        self.last_postal_code_lookup = postal_code

        try:
            suggestions = search_address_by_postal_code(postal_code)
        except RuntimeError as exc:
            messagebox.showerror("DaData", str(exc))
            return

        if not suggestions:
            messagebox.showinfo(
                "DaData",
                "DaData не нашла адрес по этому индексу. Можно заполнить адрес вручную.",
            )
            return

        if len(suggestions) == 1:
            self.fill_address_from_dadata(suggestions[0].get("data", {}))
            return

        AddressSelectionDialog(
            self.root,
            suggestions,
            on_select=self.fill_address_from_dadata,
        )

    def fill_address_from_dadata(self, address_data):
        mapping = {
            "postal_code": address_data.get("postal_code", ""),
            "region": address_data.get("region_with_type") or address_data.get("region", ""),
            "district": address_data.get("area_with_type") or address_data.get("area", ""),
            "city": address_data.get("city_with_type") or address_data.get("settlement_with_type", ""),
            "locality": address_data.get("settlement_with_type", ""),
            "street": address_data.get("street_with_type", ""),
            "house": address_data.get("house", ""),
            "building": address_data.get("block", ""),
        }

        for key, value in mapping.items():
            if value and key in self.entries:
                self.entries[key].delete(0, tk.END)
                self.entries[key].insert(0, value)
                self.entries[key].configure(bg=BG_OK)

    def on_passport_code_change(self, event):
        entry = event.widget
        formatted = format_passport_code(entry.get())

        entry.delete(0, tk.END)
        entry.insert(0, formatted)

        digits = only_digits(formatted)

        if digits:
            entry.configure(bg=BG_OK if len(digits) == 6 else BG_ERROR)
            self.update_code_results(formatted)
        else:
            entry.configure(bg=BG_OK)
            self.code_results_box.delete(0, tk.END)

    def update_code_results(self, query):
        self.code_results_box.delete(0, tk.END)
        self.current_code_results = search_passport_codes(query, limit=10)

        for code, name in self.current_code_results:
            self.code_results_box.insert(tk.END, f"{code} — {name}")

    def select_passport_code(self, event):
        selection = self.code_results_box.curselection()

        if not selection:
            return

        index = selection[0]
        code, name = self.current_code_results[index]

        self.entries["passport_code"].delete(0, tk.END)
        self.entries["passport_code"].insert(0, code)
        self.entries["passport_code"].configure(bg=BG_OK)

        self.entries["passport_issued_by"].delete(0, tk.END)
        self.entries["passport_issued_by"].insert(0, name)
        self.entries["passport_issued_by"].configure(bg=BG_OK)

    def auto_fill_genitive(self):
        surname = self.entries["surname"].get().strip()
        name = self.entries["name"].get().strip()
        patronymic = self.entries["patronymic"].get().strip()

        genitive = fio_to_genitive(surname, name, patronymic)

        self.entries["fio_genitive"].delete(0, tk.END)
        self.entries["fio_genitive"].insert(0, genitive)

    def validate_fields(self):
        errors = []

        for key, max_len in digit_limits.items():
            value = self.entries[key].get().strip()

            if value and len(value) != max_len:
                self.entries[key].configure(bg=BG_ERROR)
                errors.append(key)
            else:
                self.entries[key].configure(bg=BG_OK)

        passport_code = self.entries["passport_code"].get().strip()
        if passport_code and len(only_digits(passport_code)) != 6:
            self.entries["passport_code"].configure(bg=BG_ERROR)
            errors.append("passport_code")
        else:
            self.entries["passport_code"].configure(bg=BG_OK)

        for key in manual_date_keys:
            value = self.entries[key].get().strip()

            if value and len(value) != 10:
                self.entries[key].configure(bg=BG_ERROR)
                errors.append(key)
            else:
                self.entries[key].configure(bg=BG_OK)

        return len(errors) == 0

    def collect_data(self):
        data = {}

        for key, entry in self.entries.items():
            data[key] = entry.get().strip()

        data["passport_code"] = normalize_passport_code(data["passport_code"])
        data["fill_date"] = self.fill_date_entry.get().strip()
        data["fio"] = f"{data['surname']} {data['name']} {data['patronymic']}".strip()

        if not data["fio_genitive"]:
            data["fio_genitive"] = fio_to_genitive(
                data["surname"],
                data["name"],
                data["patronymic"]
            )

        data.update(split_date(data["fill_date"], "fill"))
        data.update(split_date(data["birth_date"], "birth"))
        data.update(split_date(data["passport_issue_date"], "passport_issue"))

        data["full_address"] = (
            f"{data['postal_code']}, {data['region']}, {data['district']}, "
            f"{data['city']}, {data['locality']}, {data['street']}, "
            f"д. {data['house']}, корп. {data['building']}, кв. {data['apartment']}"
        )

        data["passport_full"] = (
            f"{data['passport_series']} {data['passport_number']}, "
            f"выдан: {data['passport_issued_by']} {data['passport_issue_date']}, "
            f"к/п {data['passport_code']}"
        )

        creditors, total_debt = prepare_creditors_for_documents(self.creditors)
        data["creditors"] = creditors
        data["total_debt"] = total_debt
        data["total_debt_formatted"] = format_money(total_debt)

        return data

    def fill_form(self, data):
        for key, entry in self.entries.items():
            entry.delete(0, tk.END)
            entry.insert(0, data.get(key, ""))
            entry.configure(bg=BG_OK)

        if data.get("fill_date"):
            self.fill_date_entry.delete(0, tk.END)
            self.fill_date_entry.insert(0, data.get("fill_date"))

        self.creditors = data.get("creditors", [])
        self.refresh_creditors_table()

    def refresh_clients_list(self):
        self.clients = load_clients()

        names = []
        for client in self.clients:
            fio = client.get("fio", "").strip()
            if fio:
                names.append(fio)

        self.client_combo["values"] = names

    def save_client_to_base(self):
        if not self.validate_fields():
            messagebox.showerror("Ошибка", "Проверь красные поля.")
            return

        data = self.collect_data()

        if not data["surname"] or not data["name"]:
            messagebox.showerror("Ошибка", "Заполни хотя бы фамилию и имя клиента.")
            return

        self.clients = load_clients()
        updated = False

        for index, client in enumerate(self.clients):
            if client.get("fio") == data["fio"]:
                self.clients[index] = data
                updated = True
                break

        if not updated:
            self.clients.append(data)

        save_clients(self.clients)
        self.refresh_clients_list()
        messagebox.showinfo("Yurist+", "Клиент сохранён в базу.")

    def load_selected_client(self):
        selected_fio = self.client_combo.get()

        if not selected_fio:
            messagebox.showerror("Ошибка", "Выбери клиента из списка.")
            return

        self.clients = load_clients()

        for client in self.clients:
            if client.get("fio") == selected_fio:
                self.fill_form(client)
                messagebox.showinfo("Yurist+", "Клиент загружен.")
                return

    def generate_documents(self):
        if not self.validate_fields():
            messagebox.showerror("Ошибка", "Проверь красные поля.")
            return

        data = self.collect_data()
        self.save_client_to_base()

        client_folder_name = safe_folder_name(
            f"{data['surname']} {data['name']} {data['patronymic']}"
        )
        client_folder = os.path.join(OUTPUT_DIR, client_folder_name)
        os.makedirs(client_folder, exist_ok=True)

        missing_templates = [
            template_path
            for template_path in TEMPLATES.values()
            if not os.path.exists(template_path)
        ]
        if missing_templates:
            messagebox.showerror(
                "Ошибка",
                "Не найдены шаблоны:\n" + "\n".join(missing_templates),
            )
            return

        generated_files = []
        errors = []

        for doc_name, template_path in TEMPLATES.items():
            output_path = os.path.join(
                client_folder,
                f"{doc_name}_{data['surname']}_{data['name']}.docx"
            )

            try:
                doc = DocxTemplate(template_path)
                render_data = data

                if doc_name == "creditors":
                    render_data = get_creditors_template_render_data(data)

                doc.render(render_data)
                doc.save(output_path)

                if doc_name == "creditors":
                    render_creditors_table_with_python_docx(
                        output_path,
                        data.get("creditors", []),
                    )

                generated_files.append(output_path)
            except Exception as exc:
                errors.append(f"{template_path}: {exc}")

        open_folder(client_folder)

        if errors:
            messagebox.showerror(
                "Ошибка",
                "Не удалось сформировать часть документов:\n" + "\n".join(errors),
            )
            return

        messagebox.showinfo(
            "Готово",
            f"Документы сформированы: {len(generated_files)} из {len(TEMPLATES)}. "
            "Папка клиента открыта.",
        )


class GosuslugiImportDialog:
    def __init__(self, parent, on_import):
        self.parent = parent
        self.on_import = on_import

        self.window = tk.Toplevel(parent)
        self.window.title("Данные с Госуслуг")
        self.window.geometry("760x560")
        self.window.minsize(620, 420)
        self.window.transient(parent)
        self.window.grab_set()

        self.create_ui()

    def create_ui(self):
        frame = ttk.Frame(self.window, padding=10)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(1, weight=1)

        tk.Label(
            frame,
            text="Вставьте скопированный блок данных с Госуслуг:",
            font=FONT,
        ).grid(row=0, column=0, sticky="w", pady=(0, 6))

        self.text = tk.Text(frame, font=FONT, wrap="word")
        self.text.grid(row=1, column=0, sticky="nsew")

        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=self.text.yview)
        scrollbar.grid(row=1, column=1, sticky="ns")
        self.text.configure(yscrollcommand=scrollbar.set)

        buttons_frame = ttk.Frame(frame)
        buttons_frame.grid(row=2, column=0, columnspan=2, sticky="e", pady=(10, 0))

        tk.Button(
            buttons_frame,
            text="Разобрать и заполнить",
            font=FONT,
            command=self.parse_and_fill,
            width=22,
        ).pack(side="left", padx=(0, 8))

        tk.Button(
            buttons_frame,
            text="Отмена",
            font=FONT,
            command=self.window.destroy,
            width=14,
        ).pack(side="left")

    def parse_and_fill(self):
        text = self.text.get("1.0", tk.END).strip()
        if not text:
            messagebox.showerror("Госуслуги", "Вставьте текст с Госуслуг.", parent=self.window)
            return

        found_fields = self.on_import(text)
        if not found_fields:
            messagebox.showinfo(
                "Госуслуги",
                "Не удалось найти данные для заполнения.",
                parent=self.window,
            )
            return

        field_names = {
            key: label
            for label, key in fields
        }
        found_text = ", ".join(field_names.get(key, key) for key in found_fields)
        messagebox.showinfo(
            "Госуслуги",
            f"Найдено и заполнено: {found_text}",
            parent=self.window,
        )
        self.window.destroy()


class AddressSelectionDialog:
    def __init__(self, parent, suggestions, on_select):
        self.parent = parent
        self.suggestions = suggestions
        self.on_select = on_select

        self.window = tk.Toplevel(parent)
        self.window.title("Выбор адреса")
        self.window.geometry("720x320")
        self.window.minsize(560, 260)
        self.window.transient(parent)
        self.window.grab_set()

        self.create_ui()

    def create_ui(self):
        frame = ttk.Frame(self.window, padding=10)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(1, weight=1)

        tk.Label(
            frame,
            text="Найдено несколько адресов. Выберите нужный:",
            font=FONT,
        ).grid(row=0, column=0, sticky="w", pady=(0, 6))

        self.listbox = tk.Listbox(frame, font=FONT, exportselection=False)
        self.listbox.grid(row=1, column=0, sticky="nsew")
        self.listbox.bind("<Double-Button-1>", lambda event: self.select_address())

        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=self.listbox.yview)
        scrollbar.grid(row=1, column=1, sticky="ns")
        self.listbox.configure(yscrollcommand=scrollbar.set)

        for suggestion in self.suggestions:
            self.listbox.insert(tk.END, suggestion.get("value", ""))

        buttons_frame = ttk.Frame(frame)
        buttons_frame.grid(row=2, column=0, columnspan=2, sticky="e", pady=(10, 0))

        tk.Button(
            buttons_frame,
            text="Выбрать",
            font=FONT,
            command=self.select_address,
            width=14,
        ).pack(side="left", padx=(0, 8))

        tk.Button(
            buttons_frame,
            text="Отмена",
            font=FONT,
            command=self.window.destroy,
            width=14,
        ).pack(side="left")

    def select_address(self):
        selection = self.listbox.curselection()
        if not selection:
            messagebox.showerror("DaData", "Выберите адрес из списка.", parent=self.window)
            return

        suggestion = self.suggestions[selection[0]]
        self.on_select(suggestion.get("data", {}))
        self.window.destroy()


class CreditorDialog:
    def __init__(self, parent, title, on_save, creditor=None):
        self.parent = parent
        self.on_save = on_save
        self.creditor = creditor or {}
        self.bank_results = []
        self.fields = {}

        self.window = tk.Toplevel(parent)
        self.window.title(title)
        self.window.geometry("720x520")
        self.window.minsize(640, 460)
        self.window.transient(parent)
        self.window.grab_set()

        self.create_ui()
        self.fill_fields()
        self.refresh_bank_results()

    def create_ui(self):
        self.window.columnconfigure(0, weight=1)
        self.window.columnconfigure(1, weight=1)
        self.window.rowconfigure(1, weight=1)

        search_frame = ttk.Frame(self.window, padding=(10, 10, 10, 4))
        search_frame.grid(row=0, column=0, columnspan=2, sticky="ew")
        search_frame.columnconfigure(0, weight=1)

        tk.Label(search_frame, text="Поиск банка в базе", font=FONT).grid(
            row=0, column=0, sticky="w"
        )
        self.bank_search_var = tk.StringVar()
        bank_search_entry = tk.Entry(search_frame, textvariable=self.bank_search_var, font=FONT)
        bank_search_entry.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        bank_search_entry.bind("<KeyRelease>", lambda event: self.refresh_bank_results())

        results_frame = ttk.Frame(self.window, padding=(10, 0, 6, 8))
        results_frame.grid(row=1, column=0, sticky="nsew")
        results_frame.columnconfigure(0, weight=1)
        results_frame.rowconfigure(0, weight=1)

        self.bank_results_box = tk.Listbox(results_frame, height=10, font=FONT, exportselection=False)
        self.bank_results_box.grid(row=0, column=0, sticky="nsew")
        self.bank_results_box.bind("<<ListboxSelect>>", self.select_bank)
        self.bank_results_box.bind("<Double-Button-1>", self.select_bank)

        scrollbar = ttk.Scrollbar(results_frame, orient="vertical", command=self.bank_results_box.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.bank_results_box.configure(yscrollcommand=scrollbar.set)

        form_frame = ttk.Frame(self.window, padding=(6, 0, 10, 8))
        form_frame.grid(row=1, column=1, sticky="nsew")
        form_frame.columnconfigure(1, weight=1)

        for row, (label_text, key) in enumerate((
            ("Банк / кредитор", "name"),
            ("ИНН", "inn"),
            ("ОГРН", "ogrn"),
            ("Адрес", "address"),
            ("Сумма долга", "debt_sum"),
            ("Дата договора", "contract_date"),
            ("Номер договора / идентификатор", "contract_number"),
        )):
            tk.Label(form_frame, text=label_text, font=FONT).grid(
                row=row, column=0, sticky="w", padx=(0, 8), pady=4
            )
            entry = tk.Entry(form_frame, font=FONT)
            entry.grid(row=row, column=1, sticky="ew", pady=4)

            if key == "contract_date":
                entry.bind("<KeyRelease>", auto_date_format)

            self.fields[key] = entry

        buttons_frame = ttk.Frame(self.window, padding=(10, 0, 10, 10))
        buttons_frame.grid(row=2, column=0, columnspan=2, sticky="e")

        tk.Button(
            buttons_frame,
            text="Сохранить",
            font=FONT,
            command=self.save,
            width=14,
        ).pack(side="left", padx=(0, 8))

        tk.Button(
            buttons_frame,
            text="Отмена",
            font=FONT,
            command=self.window.destroy,
            width=14,
        ).pack(side="left")

    def fill_fields(self):
        for key, entry in self.fields.items():
            entry.insert(0, self.creditor.get(key, ""))

    def refresh_bank_results(self):
        self.bank_results_box.delete(0, tk.END)

        try:
            self.bank_results = search_banks(self.bank_search_var.get(), limit=50)
        except BanksDatabaseError as exc:
            self.bank_results = []
            self.bank_results_box.insert(tk.END, str(exc))
            return

        for bank in self.bank_results:
            self.bank_results_box.insert(
                tk.END,
                f"{bank.get('name', '')} | {bank.get('inn', '')} | {bank.get('ogrn', '')}",
            )

    def select_bank(self, event=None):
        selection = self.bank_results_box.curselection()
        if not selection or selection[0] >= len(self.bank_results):
            return

        bank = self.bank_results[selection[0]]
        self.set_field("name", bank.get("name", ""))
        self.set_field("inn", bank.get("inn", ""))
        self.set_field("ogrn", bank.get("ogrn", ""))
        self.set_field("address", bank.get("address", ""))

    def set_field(self, key, value):
        self.fields[key].delete(0, tk.END)
        self.fields[key].insert(0, value)

    def collect_data(self):
        return {
            "name": self.fields["name"].get().strip(),
            "inn": self.fields["inn"].get().strip(),
            "ogrn": self.fields["ogrn"].get().strip(),
            "address": self.fields["address"].get().strip(),
            "debt_sum": self.fields["debt_sum"].get().strip(),
            "contract_date": self.fields["contract_date"].get().strip(),
            "contract_number": self.fields["contract_number"].get().strip(),
        }

    def save(self):
        data = self.collect_data()
        if not data["name"]:
            messagebox.showerror("Кредиторы", "Заполните название банка или кредитора.", parent=self.window)
            return

        if self.on_save(data):
            self.window.destroy()


class BanksWindow:
    SEARCH_PLACEHOLDER = "Поиск банка по названию, ИНН или ОГРН"

    def __init__(self, parent):
        self.parent = parent
        self.window = tk.Toplevel(parent)
        self.window.title("Банки")
        self.window.geometry("900x600")
        self.window.minsize(760, 480)

        self.current_results = []
        self.search_placeholder_active = True

        self.create_ui()
        self.show_search_placeholder()
        self.refresh_results()

    def create_ui(self):
        self.window.columnconfigure(0, weight=1)
        self.window.rowconfigure(1, weight=1)

        top_frame = ttk.Frame(self.window, padding=(10, 10, 10, 6))
        top_frame.grid(row=0, column=0, sticky="ew")
        top_frame.columnconfigure(0, weight=1)

        self.search_var = tk.StringVar()
        self.search_entry = tk.Entry(
            top_frame,
            textvariable=self.search_var,
            font=("Times New Roman", 14),
        )
        self.search_entry.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        self.search_entry.bind("<FocusIn>", self.on_search_focus_in)
        self.search_entry.bind("<FocusOut>", self.on_search_focus_out)
        self.search_entry.bind("<KeyRelease>", self.on_search_change)

        tk.Button(
            top_frame,
            text="Добавить банк",
            font=FONT,
            command=self.open_add_dialog,
            width=18,
        ).grid(row=0, column=1, sticky="e")

        list_frame = ttk.Frame(self.window, padding=(10, 0, 10, 10))
        list_frame.grid(row=1, column=0, sticky="nsew")
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)

        self.banks_tree = ttk.Treeview(
            list_frame,
            columns=("bank",),
            show="headings",
            selectmode="browse",
        )
        self.banks_tree.heading("bank", text="Название банка | ИНН | ОГРН")
        self.banks_tree.column("bank", width=820, minwidth=500, stretch=True)
        self.banks_tree.grid(row=0, column=0, sticky="nsew")
        self.banks_tree.bind("<Double-Button-1>", self.open_details_for_selected)
        self.banks_tree.bind("<Button-3>", self.show_context_menu)
        self.banks_tree.bind("<Button-2>", self.show_context_menu)

        scrollbar = ttk.Scrollbar(
            list_frame,
            orient="vertical",
            command=self.banks_tree.yview,
        )
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.banks_tree.configure(yscrollcommand=scrollbar.set)

        self.context_menu = tk.Menu(self.window, tearoff=0)
        self.context_menu.add_command(
            label="Изменить",
            command=self.open_edit_dialog_for_selected,
        )
        self.context_menu.add_command(label="Удалить", command=self.remove_selected_bank)

    def show_search_placeholder(self):
        self.search_placeholder_active = True
        self.search_var.set(self.SEARCH_PLACEHOLDER)
        self.search_entry.configure(fg="gray")

    def hide_search_placeholder(self):
        if self.search_placeholder_active:
            self.search_placeholder_active = False
            self.search_var.set("")
            self.search_entry.configure(fg="black")

    def on_search_focus_in(self, event):
        self.hide_search_placeholder()

    def on_search_focus_out(self, event):
        if not self.search_var.get().strip():
            self.show_search_placeholder()

    def on_search_change(self, event):
        if not self.search_placeholder_active:
            self.refresh_results()

    def get_search_query(self):
        if self.search_placeholder_active:
            return ""

        return self.search_var.get().strip()

    def refresh_results(self):
        for item_id in self.banks_tree.get_children():
            self.banks_tree.delete(item_id)

        try:
            self.current_results = search_banks(self.get_search_query())
        except BanksDatabaseError as exc:
            self.current_results = []
            messagebox.showerror("База банков", str(exc), parent=self.window)
            return

        for bank in self.current_results:
            self.banks_tree.insert(
                "",
                tk.END,
                iid=str(bank["id"]),
                values=(self.format_bank_row(bank),),
            )

    def format_bank_row(self, bank):
        return " | ".join((
            bank.get("name", ""),
            bank.get("inn", ""),
            bank.get("ogrn", ""),
        ))

    def get_selected_bank_id(self):
        selection = self.banks_tree.selection()
        if not selection:
            return None

        return int(selection[0])

    def select_row_under_pointer(self, event):
        row_id = self.banks_tree.identify_row(event.y)
        if row_id:
            self.banks_tree.selection_set(row_id)
            self.banks_tree.focus(row_id)
            return True

        return False

    def show_context_menu(self, event):
        if not self.select_row_under_pointer(event):
            return

        try:
            self.context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.context_menu.grab_release()

    def open_add_dialog(self):
        BankEditDialog(
            self.window,
            title="Добавить банк",
            on_save=self.add_bank_from_dialog,
        )

    def open_edit_dialog_for_selected(self):
        bank_id = self.get_selected_bank_id()
        if not bank_id:
            messagebox.showerror("База банков", "Выберите банк для изменения.", parent=self.window)
            return

        try:
            bank = get_bank_details(bank_id)
        except BanksDatabaseError as exc:
            messagebox.showerror("База банков", str(exc), parent=self.window)
            return

        if not bank:
            messagebox.showerror("База банков", "Банк не найден.", parent=self.window)
            self.refresh_results()
            return

        BankEditDialog(
            self.window,
            title="Изменить банк",
            bank=bank,
            on_save=lambda data: self.update_bank_from_dialog(bank_id, data),
        )

    def open_details_for_selected(self, event=None):
        bank_id = self.get_selected_bank_id()
        if not bank_id:
            return

        try:
            bank = get_bank_details(bank_id)
        except BanksDatabaseError as exc:
            messagebox.showerror("База банков", str(exc), parent=self.window)
            return

        if not bank:
            messagebox.showerror("База банков", "Банк не найден.", parent=self.window)
            self.refresh_results()
            return

        BankDetailsDialog(self.window, bank, on_edit=self.open_edit_dialog_for_selected)

    def add_bank_from_dialog(self, data):
        try:
            add_bank(data)
        except BanksDatabaseError as exc:
            messagebox.showerror("База банков", str(exc), parent=self.window)
            return False

        self.refresh_results()
        messagebox.showinfo("База банков", "Банк добавлен.", parent=self.window)
        return True

    def update_bank_from_dialog(self, bank_id, data):
        try:
            update_bank(bank_id, data)
        except BanksDatabaseError as exc:
            messagebox.showerror("База банков", str(exc), parent=self.window)
            return False

        self.refresh_results()
        tree_id = str(bank_id)
        if self.banks_tree.exists(tree_id):
            self.banks_tree.selection_set(tree_id)
            self.banks_tree.focus(tree_id)
        messagebox.showinfo("База банков", "Изменения банка сохранены.", parent=self.window)
        return True

    def remove_selected_bank(self):
        bank_id = self.get_selected_bank_id()
        if not bank_id:
            messagebox.showerror("База банков", "Выберите банк для удаления.", parent=self.window)
            return

        bank_name = self.banks_tree.item(str(bank_id), "values")[0].split(" | ", 1)[0]
        if not messagebox.askyesno(
            "Удаление банка",
            f"Удалить банк «{bank_name}» из базы данных?",
            parent=self.window,
        ):
            return

        try:
            delete_bank(bank_id)
        except BanksDatabaseError as exc:
            messagebox.showerror("База банков", str(exc), parent=self.window)
            return

        self.refresh_results()
        messagebox.showinfo("База банков", "Банк удалён.", parent=self.window)


class BankEditDialog:
    def __init__(self, parent, title, on_save, bank=None):
        self.parent = parent
        self.on_save = on_save
        self.bank = bank or {}
        self.fields = {}

        self.window = tk.Toplevel(parent)
        self.window.title(title)
        self.window.geometry("520x270")
        self.window.resizable(False, False)
        self.window.transient(parent)
        self.window.grab_set()

        self.create_ui()
        self.fill_fields()
        self.fields["name"].focus_set()

    def create_ui(self):
        form_frame = ttk.Frame(self.window, padding=12)
        form_frame.pack(fill="both", expand=True)
        form_frame.columnconfigure(1, weight=1)

        for row, (label_text, key) in enumerate((
            ("Название", "name"),
            ("ИНН", "inn"),
            ("ОГРН", "ogrn"),
            ("Адрес", "address"),
        )):
            tk.Label(form_frame, text=label_text, font=FONT).grid(
                row=row, column=0, sticky="w", padx=(0, 8), pady=6
            )
            entry = tk.Entry(form_frame, font=FONT)
            entry.grid(row=row, column=1, sticky="ew", pady=6)
            self.fields[key] = entry

        buttons_frame = ttk.Frame(form_frame)
        buttons_frame.grid(row=4, column=0, columnspan=2, sticky="e", pady=(14, 0))

        tk.Button(
            buttons_frame,
            text="Сохранить",
            font=FONT,
            command=self.save,
            width=14,
        ).pack(side="left", padx=(0, 8))

        tk.Button(
            buttons_frame,
            text="Отмена",
            font=FONT,
            command=self.window.destroy,
            width=14,
        ).pack(side="left")

    def fill_fields(self):
        for key, entry in self.fields.items():
            entry.insert(0, self.bank.get(key, ""))

    def collect_data(self):
        return {
            "name": self.fields["name"].get().strip(),
            "inn": self.fields["inn"].get().strip(),
            "ogrn": self.fields["ogrn"].get().strip(),
            "address": self.fields["address"].get().strip(),
        }

    def save(self):
        if self.on_save(self.collect_data()):
            self.window.destroy()


class BankDetailsDialog:
    def __init__(self, parent, bank, on_edit=None):
        self.parent = parent
        self.bank = bank
        self.on_edit = on_edit

        self.window = tk.Toplevel(parent)
        self.window.title("Информация о банке")
        self.window.geometry("560x420")
        self.window.minsize(460, 340)
        self.window.transient(parent)

        self.create_ui()

    def create_ui(self):
        frame = ttk.Frame(self.window, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)
        frame.rowconfigure(4, weight=1)
        frame.rowconfigure(5, weight=1)

        self.add_value_row(frame, "Название", self.bank.get("name", ""), 0)
        self.add_value_row(frame, "ИНН", self.bank.get("inn", ""), 1)
        self.add_value_row(frame, "ОГРН", self.bank.get("ogrn", ""), 2)
        self.add_value_row(frame, "Адрес", self.bank.get("address", ""), 3)
        self.add_text_row(frame, "Филиалы", "\n".join(self.bank.get("branches", [])), 4)
        self.add_text_row(frame, "Альтернативные названия", "\n".join(self.bank.get("aliases", [])), 5)

        buttons_frame = ttk.Frame(frame)
        buttons_frame.grid(row=6, column=0, columnspan=2, sticky="e", pady=(10, 0))

        if self.on_edit:
            tk.Button(
                buttons_frame,
                text="Изменить",
                font=FONT,
                command=self.open_edit,
                width=14,
            ).pack(side="left", padx=(0, 8))

        tk.Button(
            buttons_frame,
            text="Закрыть",
            font=FONT,
            command=self.window.destroy,
            width=14,
        ).pack(side="left")

    def add_value_row(self, frame, label_text, value, row):
        tk.Label(frame, text=label_text, font=FONT).grid(
            row=row, column=0, sticky="nw", padx=(0, 8), pady=4
        )
        tk.Label(
            frame,
            text=value or "—",
            font=FONT,
            anchor="w",
            justify="left",
            wraplength=360,
        ).grid(row=row, column=1, sticky="ew", pady=4)

    def add_text_row(self, frame, label_text, value, row):
        tk.Label(frame, text=label_text, font=FONT).grid(
            row=row, column=0, sticky="nw", padx=(0, 8), pady=4
        )
        text = tk.Text(frame, height=4, font=FONT, wrap="word")
        text.grid(row=row, column=1, sticky="nsew", pady=4)
        text.insert("1.0", value or "—")
        text.configure(state="disabled")

    def open_edit(self):
        self.window.destroy()
        self.on_edit()


if __name__ == "__main__":
    ensure_files()
    create_passport_codes_db_if_needed()
    root = tk.Tk()
    app = YuristApp(root)
    root.mainloop()